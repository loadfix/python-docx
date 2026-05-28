"""Unit-test suite for ``docx.kit.medical`` template factories."""

from __future__ import annotations

from io import BytesIO

import pytest

from docx.document import Document as DocumentCls
from docx.kit import medical


# -- Shared fixtures ------------------------------------------------------


@pytest.fixture
def patient():
    return {
        "name": "John Doe",
        "dob": "1980-05-15",
        "mrn": "1234567",
        "medicare": "2956 12345 1",
        "sex": "M",
    }


@pytest.fixture
def provider():
    return {
        "name": "Dr. Alice Smith",
        "role": "GP",
        "practice": "Sydney Medical Centre",
        "provider_number": "1234567A",
    }


@pytest.fixture
def vitals():
    return {
        "bp": "120/80",
        "hr": 72,
        "temp": 36.8,
        "rr": 18,
        "spo2": 98,
    }


def _texts(document: DocumentCls):
    """Return the text of every paragraph in ``document``."""
    return [p.text for p in document.paragraphs]


def _full_text(document: DocumentCls) -> str:
    return "\n".join(_texts(document))


def _all_cell_texts(document: DocumentCls):
    out = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                out.append(cell.text)
    return out


# -- SOAP note ------------------------------------------------------------


class DescribeSoapNote:
    """Unit-test suite for ``medical.soap_note``."""

    def it_returns_a_document_with_a_SOAP_title(self, patient, provider):
        doc = medical.soap_note(patient=patient, provider=provider)

        assert isinstance(doc, DocumentCls)
        assert "Clinical Note (SOAP)" in _full_text(doc)

    def it_includes_the_four_SOAP_sections(self, patient, provider):
        doc = medical.soap_note(patient=patient, provider=provider)
        text = _full_text(doc)

        # -- Subjective / Objective / Assessment / Plan are the four
        # -- canonical SOAP headings the issue acceptance calls out. --
        assert "Subjective" in text
        assert "Objective" in text
        assert "Assessment" in text
        assert "Plan" in text

    def it_includes_the_template_disclaimer(self, patient, provider):
        doc = medical.soap_note(patient=patient, provider=provider)
        text = _full_text(doc)

        assert "TEMPLATE ONLY" in text
        assert "not itself a medical record" in text
        assert "not legal advice" in text

    def it_renders_the_patient_block_as_a_table(self, patient, provider):
        doc = medical.soap_note(patient=patient, provider=provider)

        cells = _all_cell_texts(doc)
        assert "Name" in cells
        assert "John Doe" in cells
        assert "Date of birth" in cells
        assert "1980-05-15" in cells
        assert "MRN" in cells
        assert "1234567" in cells

    def it_renders_unrecognised_patient_keys_verbatim(self, provider):
        doc = medical.soap_note(
            patient={"name": "John Doe", "next_of_kin": "Jane Doe"},
            provider=provider,
        )

        cells = _all_cell_texts(doc)
        assert "Next Of Kin" in cells
        assert "Jane Doe" in cells

    def it_renders_the_provider_details(self, patient, provider):
        doc = medical.soap_note(patient=patient, provider=provider)
        text = _full_text(doc)

        assert "Dr. Alice Smith" in text
        assert "GP" in text
        assert "Sydney Medical Centre" in text
        assert "Provider No: 1234567A" in text

    def it_renders_the_encounter_date(self, patient, provider):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            encounter_date="2026-03-15",
        )

        assert "Encounter Date: 2026-03-15" in _full_text(doc)

    def it_renders_subjective_as_a_string(self, patient, provider):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            subjective="Patient reports persistent cough for 3 weeks.",
        )

        assert (
            "Patient reports persistent cough for 3 weeks."
            in _full_text(doc)
        )

    def it_renders_subjective_as_a_sequence(self, patient, provider):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            subjective=["Cough for 3 weeks.", "No fever or chills."],
        )
        text = _full_text(doc)

        assert "Cough for 3 weeks." in text
        assert "No fever or chills." in text

    def it_renders_vitals_as_a_structured_table(
        self, patient, provider, vitals
    ):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            objective={"vitals": vitals},
        )

        cells = _all_cell_texts(doc)
        # -- BP/HR/Temp/RR/SpO2 must each appear in the vitals table --
        assert "Blood pressure" in cells
        assert "120/80 mmHg" in cells
        assert "Heart rate" in cells
        assert "72 bpm" in cells
        assert "Temperature" in cells
        assert "36.8 °C" in cells
        assert "Respiratory rate" in cells
        assert "SpO2" in cells

    def it_renders_unrecognised_vitals_keys_verbatim(self, patient, provider):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            objective={"vitals": {"bp": "110/70", "fio2": "21%"}},
        )

        cells = _all_cell_texts(doc)
        assert "Fio2" in cells
        assert "21%" in cells

    def it_renders_examination_text(self, patient, provider):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            objective={
                "examination": "Chest clear, no rales, no wheeze.",
            },
        )

        assert "Chest clear, no rales, no wheeze." in _full_text(doc)

    def it_renders_labs_as_a_bulleted_list(self, patient, provider):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            objective={"labs": ["FBC: WNL", "CRP: 12 mg/L"]},
        )
        text = _full_text(doc)

        assert "Investigations" in text
        assert "FBC: WNL" in text
        assert "CRP: 12 mg/L" in text

    def it_renders_assessment_with_string_diagnoses(self, patient, provider):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            assessment=["Acute bronchitis (J20.9)", "Hypertension stable"],
        )
        text = _full_text(doc)

        assert "Acute bronchitis (J20.9)" in text
        assert "Hypertension stable" in text

    def it_renders_assessment_with_coded_diagnoses(self, patient, provider):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            assessment=[
                {"text": "Acute bronchitis", "code": "J20.9"},
                {"text": "Hypertension"},
            ],
        )
        text = _full_text(doc)

        assert "Acute bronchitis (J20.9)" in text
        assert "Hypertension" in text

    def it_renders_plan_as_a_numbered_list(self, patient, provider):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            plan=[
                "Amoxicillin 500mg TDS for 7 days",
                "Review in 1 week",
                "Continue current antihypertensive regimen",
            ],
        )
        text = _full_text(doc)

        assert "Amoxicillin 500mg TDS for 7 days" in text
        assert "Review in 1 week" in text
        assert "Continue current antihypertensive regimen" in text

    def it_renders_placeholders_when_sections_missing(
        self, patient, provider
    ):
        doc = medical.soap_note(patient=patient, provider=provider)
        text = _full_text(doc)

        assert "[Insert patient-reported history" in text
        assert "[Insert clinical impression" in text
        assert "[Insert management plan" in text

    def it_includes_a_signature_block(self, patient, provider):
        doc = medical.soap_note(patient=patient, provider=provider)
        text = _full_text(doc)

        assert "Signed" in text
        assert "Clinician: Dr. Alice Smith" in text
        assert "Date: ______________________________" in text

    def it_raises_when_patient_missing(self, provider):
        with pytest.raises(ValueError, match="patient is required"):
            medical.soap_note(provider=provider)

    def it_raises_when_patient_has_no_name(self, provider):
        with pytest.raises(ValueError, match="non-empty 'name'"):
            medical.soap_note(
                patient={"dob": "1980-05-15"}, provider=provider
            )

    def it_raises_when_provider_missing(self, patient):
        with pytest.raises(ValueError, match="provider is required"):
            medical.soap_note(patient=patient)

    def it_raises_when_provider_has_no_name(self, patient):
        with pytest.raises(ValueError, match="non-empty 'name'"):
            medical.soap_note(
                patient=patient, provider={"role": "GP"}
            )


# -- Discharge summary ----------------------------------------------------


class DescribeDischargeSummary:
    """Unit-test suite for ``medical.discharge_summary``."""

    def it_returns_a_document_with_a_discharge_summary_title(
        self, patient, provider
    ):
        doc = medical.discharge_summary(patient=patient, provider=provider)

        assert isinstance(doc, DocumentCls)
        assert "Hospital Discharge Summary" in _full_text(doc)

    def it_renders_admission_and_discharge_dates(self, patient, provider):
        doc = medical.discharge_summary(
            patient=patient,
            provider=provider,
            admission_date="2026-03-10",
            discharge_date="2026-03-15",
        )
        text = _full_text(doc)

        assert "Admission: 2026-03-10" in text
        assert "Discharge: 2026-03-15" in text

    def it_renders_only_admission_date_when_discharge_missing(
        self, patient, provider
    ):
        doc = medical.discharge_summary(
            patient=patient,
            provider=provider,
            admission_date="2026-03-10",
        )

        assert "Admission: 2026-03-10" in _full_text(doc)

    def it_renders_the_presenting_complaint(self, patient, provider):
        doc = medical.discharge_summary(
            patient=patient,
            provider=provider,
            presenting_complaint="Acute chest pain, ED presentation.",
        )
        text = _full_text(doc)

        assert "Presenting Complaint" in text
        assert "Acute chest pain, ED presentation." in text

    def it_renders_diagnoses_as_a_numbered_list(self, patient, provider):
        doc = medical.discharge_summary(
            patient=patient,
            provider=provider,
            diagnoses=[
                {"text": "Non-ST-elevation MI", "code": "I21.4"},
                "Hyperlipidaemia",
            ],
        )
        text = _full_text(doc)

        assert "Final Diagnoses" in text
        assert "Non-ST-elevation MI (I21.4)" in text
        assert "Hyperlipidaemia" in text

    def it_renders_discharge_medications(self, patient, provider):
        doc = medical.discharge_summary(
            patient=patient,
            provider=provider,
            discharge_medications=[
                {
                    "name": "Aspirin",
                    "dose": "100mg",
                    "frequency": "mane",
                    "duration": "ongoing",
                },
                "Atorvastatin 40mg nocte",
            ],
        )
        text = _full_text(doc)

        assert "Discharge Medications" in text
        # -- structured form composes name + dose + freq + duration --
        assert "Aspirin 100mg mane" in text
        assert "ongoing" in text
        assert "Atorvastatin 40mg nocte" in text

    def it_renders_investigations_as_bullets(self, patient, provider):
        doc = medical.discharge_summary(
            patient=patient,
            provider=provider,
            investigations=[
                {"name": "Troponin", "result": "0.45 ng/mL (raised)"},
                "ECG: ST depression V4-V6",
            ],
        )
        text = _full_text(doc)

        assert "Investigations" in text
        assert "Troponin: 0.45 ng/mL (raised)" in text
        assert "ECG: ST depression V4-V6" in text

    def it_renders_procedures_when_supplied(self, patient, provider):
        doc = medical.discharge_summary(
            patient=patient,
            provider=provider,
            procedures=[
                {"name": "Coronary angiography", "date": "2026-03-12"},
                "Echocardiogram",
            ],
        )
        text = _full_text(doc)

        assert "Procedures" in text
        assert "Coronary angiography (2026-03-12)" in text
        assert "Echocardiogram" in text

    def it_renders_followup_plan(self, patient, provider):
        doc = medical.discharge_summary(
            patient=patient,
            provider=provider,
            follow_up=[
                "GP review in 1 week",
                "Cardiology outpatient in 6 weeks",
            ],
        )
        text = _full_text(doc)

        assert "Follow-up Plan" in text
        assert "GP review in 1 week" in text
        assert "Cardiology outpatient in 6 weeks" in text

    def it_renders_discharge_vitals_as_a_table(
        self, patient, provider, vitals
    ):
        doc = medical.discharge_summary(
            patient=patient,
            provider=provider,
            discharge_vitals=vitals,
        )

        cells = _all_cell_texts(doc)
        assert "Vitals at Discharge" in _full_text(doc)
        assert "Blood pressure" in cells
        assert "120/80 mmHg" in cells

    def it_includes_the_template_disclaimer(self, patient, provider):
        doc = medical.discharge_summary(patient=patient, provider=provider)

        assert "TEMPLATE ONLY" in _full_text(doc)

    def it_renders_placeholders_when_sections_missing(
        self, patient, provider
    ):
        doc = medical.discharge_summary(patient=patient, provider=provider)
        text = _full_text(doc)

        assert "[Insert reason for admission" in text
        assert "[Insert pathology" in text
        assert "[Insert principal" in text
        assert "[Insert discharge medication" in text
        assert "[Insert GP review" in text

    def it_raises_when_patient_missing(self, provider):
        with pytest.raises(ValueError, match="patient is required"):
            medical.discharge_summary(provider=provider)

    def it_raises_when_provider_missing(self, patient):
        with pytest.raises(ValueError, match="provider is required"):
            medical.discharge_summary(patient=patient)


# -- Referral letter ------------------------------------------------------


class DescribeReferralLetter:
    """Unit-test suite for ``medical.referral_letter``."""

    def it_returns_a_document_with_a_referral_letter_title(
        self, patient, provider
    ):
        doc = medical.referral_letter(patient=patient, referrer=provider)

        assert isinstance(doc, DocumentCls)
        assert "Referral Letter" in _full_text(doc)

    def it_renders_the_referrer_block(self, patient, provider):
        doc = medical.referral_letter(patient=patient, referrer=provider)
        text = _full_text(doc)

        assert "Dr. Alice Smith" in text
        assert "GP" in text
        assert "Sydney Medical Centre" in text

    def it_renders_a_named_recipient_with_salutation(self, patient, provider):
        doc = medical.referral_letter(
            patient=patient,
            referrer=provider,
            recipient={
                "name": "Dr. Bob Jones",
                "role": "Cardiologist",
                "practice": "Royal Prince Alfred Hospital",
                "address": "50 Missenden Rd, Camperdown NSW",
            },
        )
        text = _full_text(doc)

        assert "Dr. Bob Jones" in text
        assert "Cardiologist" in text
        assert "Royal Prince Alfred Hospital" in text
        assert "Dear Dr. Bob Jones," in text

    def it_falls_back_to_dear_colleague_without_recipient(
        self, patient, provider
    ):
        doc = medical.referral_letter(patient=patient, referrer=provider)

        assert "Dear Colleague," in _full_text(doc)

    def it_renders_the_reason_for_referral(self, patient, provider):
        doc = medical.referral_letter(
            patient=patient,
            referrer=provider,
            reason="Please assess for ischaemic heart disease.",
        )
        text = _full_text(doc)

        assert "Reason for Referral" in text
        assert "Please assess for ischaemic heart disease." in text

    def it_renders_history_and_examination(self, patient, provider):
        doc = medical.referral_letter(
            patient=patient,
            referrer=provider,
            history="3-month history of exertional chest tightness.",
            examination="BP 145/90, otherwise unremarkable cardiovascular exam.",
        )
        text = _full_text(doc)

        assert "History" in text
        assert "exertional chest tightness" in text
        assert "Examination" in text
        assert "BP 145/90" in text

    def it_renders_investigations_to_date(self, patient, provider):
        doc = medical.referral_letter(
            patient=patient,
            referrer=provider,
            investigations=[
                {"name": "Resting ECG", "result": "normal sinus rhythm"},
                "Lipid panel: LDL 4.2 mmol/L",
            ],
        )
        text = _full_text(doc)

        assert "Investigations to Date" in text
        assert "Resting ECG: normal sinus rhythm" in text
        assert "Lipid panel: LDL 4.2 mmol/L" in text

    def it_renders_current_medications(self, patient, provider):
        doc = medical.referral_letter(
            patient=patient,
            referrer=provider,
            medications=[
                {"name": "Perindopril", "dose": "5mg", "frequency": "mane"},
                "Atorvastatin 40mg nocte",
            ],
        )
        text = _full_text(doc)

        assert "Current Medications" in text
        assert "Perindopril 5mg mane" in text
        assert "Atorvastatin 40mg nocte" in text

    def it_renders_allergies(self, patient, provider):
        doc = medical.referral_letter(
            patient=patient,
            referrer=provider,
            allergies=["Penicillin (rash)", "Sulfonamides"],
        )
        text = _full_text(doc)

        assert "Allergies" in text
        assert "Penicillin (rash)" in text
        assert "Sulfonamides" in text

    def it_renders_requested_action_as_a_string(self, patient, provider):
        doc = medical.referral_letter(
            patient=patient,
            referrer=provider,
            requested_action="Please arrange exercise stress echo.",
        )
        text = _full_text(doc)

        assert "Requested Action" in text
        assert "Please arrange exercise stress echo." in text

    def it_renders_requested_action_as_a_numbered_list(
        self, patient, provider
    ):
        doc = medical.referral_letter(
            patient=patient,
            referrer=provider,
            requested_action=[
                "Assess for ischaemic heart disease",
                "Advise on statin uptitration",
            ],
        )
        text = _full_text(doc)

        assert "Assess for ischaemic heart disease" in text
        assert "Advise on statin uptitration" in text

    def it_renders_the_default_closing(self, patient, provider):
        doc = medical.referral_letter(patient=patient, referrer=provider)
        text = _full_text(doc)

        assert "Thank you for seeing this patient." in text
        assert "Kind regards," in text

    def it_renders_a_custom_closing_when_supplied(self, patient, provider):
        doc = medical.referral_letter(
            patient=patient,
            referrer=provider,
            closing="Looking forward to your assessment.",
        )

        assert "Looking forward to your assessment." in _full_text(doc)

    def it_includes_the_template_disclaimer(self, patient, provider):
        doc = medical.referral_letter(patient=patient, referrer=provider)

        assert "TEMPLATE ONLY" in _full_text(doc)

    def it_includes_a_signature_block(self, patient, provider):
        doc = medical.referral_letter(patient=patient, referrer=provider)
        text = _full_text(doc)

        assert "Signed" in text
        assert "Clinician: Dr. Alice Smith" in text

    def it_raises_when_patient_missing(self, provider):
        with pytest.raises(ValueError, match="patient is required"):
            medical.referral_letter(referrer=provider)

    def it_raises_when_referrer_missing(self, patient):
        with pytest.raises(ValueError, match="provider is required"):
            medical.referral_letter(patient=patient)


# -- Diagnosis formatter --------------------------------------------------


class DescribeFormatDiagnosis:
    """Unit-test suite for the private ``_format_diagnosis`` helper."""

    def it_returns_a_string_verbatim(self):
        assert (
            medical._format_diagnosis("Acute bronchitis")
            == "Acute bronchitis"
        )

    def it_formats_a_text_and_code_mapping(self):
        out = medical._format_diagnosis(
            {"text": "Acute bronchitis", "code": "J20.9"}
        )
        assert out == "Acute bronchitis (J20.9)"

    def it_formats_a_text_only_mapping(self):
        assert (
            medical._format_diagnosis({"text": "Hypertension"})
            == "Hypertension"
        )

    def it_formats_a_code_only_mapping(self):
        assert medical._format_diagnosis({"code": "I10"}) == "I10"

    def it_raises_on_an_empty_mapping(self):
        with pytest.raises(ValueError, match="must include 'text' or 'code'"):
            medical._format_diagnosis({})


# -- Round-trip integration ----------------------------------------------


class DescribeMedicalRoundTrip:
    """End-to-end smoke-tests: every factory produces a saveable document."""

    def it_can_save_a_SOAP_note_to_a_BytesIO(
        self, patient, provider, vitals
    ):
        doc = medical.soap_note(
            patient=patient,
            provider=provider,
            encounter_date="2026-03-15",
            subjective="Persistent cough for 3 weeks.",
            objective={
                "vitals": vitals,
                "examination": "Chest clear, no rales.",
                "labs": ["FBC: WNL", "CRP: 12 mg/L"],
            },
            assessment=["Acute bronchitis (J20.9)", "Hypertension stable"],
            plan=[
                "Amoxicillin 500mg TDS for 7 days",
                "Review in 1 week",
            ],
        )
        buf = BytesIO()
        doc.save(buf)
        # -- Word .docx is a zip; the magic bytes are 'PK' --
        assert buf.getvalue()[:2] == b"PK"

    def it_can_save_a_discharge_summary_to_a_BytesIO(
        self, patient, provider, vitals
    ):
        doc = medical.discharge_summary(
            patient=patient,
            provider=provider,
            admission_date="2026-03-10",
            discharge_date="2026-03-15",
            presenting_complaint="Acute chest pain.",
            history="3-day history of central chest pain on exertion.",
            investigations=[
                {"name": "Troponin", "result": "0.45 ng/mL"},
            ],
            procedures=[
                {"name": "Coronary angiography", "date": "2026-03-12"},
            ],
            diagnoses=[{"text": "NSTEMI", "code": "I21.4"}],
            discharge_medications=[
                {
                    "name": "Aspirin",
                    "dose": "100mg",
                    "frequency": "mane",
                },
            ],
            follow_up=["GP review in 1 week"],
            discharge_vitals=vitals,
        )
        buf = BytesIO()
        doc.save(buf)
        assert buf.getvalue()[:2] == b"PK"

    def it_can_save_a_referral_letter_to_a_BytesIO(self, patient, provider):
        doc = medical.referral_letter(
            patient=patient,
            referrer=provider,
            recipient={
                "name": "Dr. Bob Jones",
                "role": "Cardiologist",
                "practice": "RPA",
            },
            encounter_date="2026-03-15",
            reason="Please assess for ischaemic heart disease.",
            history="3-month exertional chest tightness.",
            examination="BP 145/90, otherwise unremarkable.",
            investigations=[
                {"name": "Resting ECG", "result": "normal"},
            ],
            medications=[
                {"name": "Perindopril", "dose": "5mg", "frequency": "mane"},
            ],
            allergies=["Nil known"],
            requested_action=[
                "Assess for IHD",
                "Advise on statin uptitration",
            ],
        )
        buf = BytesIO()
        doc.save(buf)
        assert buf.getvalue()[:2] == b"PK"
