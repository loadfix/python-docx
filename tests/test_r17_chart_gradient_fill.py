# pyright: reportPrivateUsage=false

"""R17-3 smoke test: chart gradient-fill authoring via ooxml-chart 0.5.

``python-ooxml-chart`` 0.5 added typed gradient-fill accessors
(``FormatFill.gradient``, ``FormatFill.apply_gradient``,
``GradientFill``, ``FILL_TYPE``, ``XL_GRADIENT_FILL_TYPE``). This
test pins that docx's ``ChartSeries.format.fill.apply_gradient(...)``
surfaces the shared proxy and that the resulting ``a:gradFill``
survives a ``Document.save`` / reopen round-trip.

Registration of the shared ``a:gradFill`` / ``a:gsLst`` / ``a:gs`` /
``a:lin`` ``CT_*`` classes in docx's ``element_class_lookup`` (in
``docx.oxml.__init__``) is what lets the shared ``GradientFill``
proxy operate against docx-parsed chart parts after reload.
"""

from __future__ import annotations

import io

from docx import Document
from docx.chart import WD_CHART_TYPE


class DescribeR17ChartGradientFill:
    """``ChartSeries.format.fill.apply_gradient`` round-trips via ``Document.save``."""

    def it_roundtrips_a_linear_gradient_on_a_series_fill(self):
        from ooxml_chart import FILL_TYPE

        document = Document()
        chart = document.add_chart(
            WD_CHART_TYPE.BAR, ["a", "b", "c"], {"S1": [1.0, 2.0, 3.0]}
        )
        series = chart.series[0]
        series.format.fill.apply_gradient(
            stops=[(0.0, "FF0000"), (1.0, "0000FF")],
            angle=45.0,
        )

        buf = io.BytesIO()
        document.save(buf)
        buf.seek(0)
        reopened = Document(buf)

        chart2 = reopened.charts[0]
        series2 = chart2.series[0]
        fill2 = series2.format.fill
        assert fill2.type is FILL_TYPE.GRADIENT
        grad = fill2.gradient
        assert grad is not None
        assert grad.stops == [(0.0, "FF0000"), (1.0, "0000FF")]
        assert grad.angle == 45.0
